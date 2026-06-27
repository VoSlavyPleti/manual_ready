"""Evaluate discrepancy-analysis artifacts against filtered legal gold rows.

The evaluator is intentionally separate from the agent. It filters definitions,
headings, requisites, signatures, blank technical rows, and non-informative
parents before scoring legal metrics.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DOC_CONFIGS = {
    "KAVKAZ": {"start": 3, "cid": 0, "desc": 1, "status": 2, "analog": 3},
    "KUZBAS": {"start": 2, "cid": 0, "desc": 1, "analog": 2, "status": 4},
    "KALUGA": {"start": 5, "cid": 1, "desc": 2, "analog": 3, "status": 4},
}

ID_RE = re.compile(r"\d+(?:\.\d+)+\.?")


def norm_id(value: Any) -> str:
    text = str(value or "").replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text).strip(" ;,")
    while text.endswith("."):
        text = text[:-1].strip()
    return text


def parse_ids(value: Any) -> set[str]:
    return {norm_id(match.group(0)) for match in ID_RE.finditer(str(value or ""))}


def normalize_status(value: Any, has_analogue: bool) -> str | None:
    text = str(value or "").lower()
    if "расхожд" in text or "partial" in text or "deviation" in text:
        return "deviation"
    if "соответ" in text or "aligned" in text or "full" in text:
        return "aligned"
    if "нет аналога" in text:
        return "extra_in_contract" if not has_analogue else "missing_in_contract"
    if has_analogue:
        return "aligned"
    return None


def numeric_key(value: str) -> tuple[int, ...] | None:
    value = norm_id(value)
    if not re.fullmatch(r"\d+(?:\.\d+)*", value):
        return None
    return tuple(int(part) for part in value.split("."))


def is_prefix(parent: str, child: str) -> bool:
    return child.startswith(parent + ".")


def looks_non_legal_description(description: str) -> bool:
    text = description.lower().strip()
    if not text:
        return True
    prefixes = (
        "определение",
        "раздел ",
        "заголовок",
        "реквизит",
        "подпись",
        "преамбула",
        "пустая форма",
        "форма без заполнения",
    )
    return text.startswith(prefixes)


def looks_legal_description(description: str) -> bool:
    text = description.lower()
    if looks_non_legal_description(description):
        return False
    legal_markers = (
        "обяз",
        "прав",
        "ответствен",
        "штраф",
        "пен",
        "неустой",
        "срок",
        "оплат",
        "прием",
        "приём",
        "растор",
        "отказ",
        "удерж",
        "контрол",
        "уведом",
        "еис",
        "44",
        "223",
        "цена",
        "мцк",
        "акт",
        "документ",
        "конфиденц",
        "персональ",
        "форс",
        "спор",
        "подсуд",
        "запрет",
        "переход",
    )
    return any(marker in text for marker in legal_markers)


def load_gold(doc: str, gold_path: Path) -> dict[str, list[dict[str, Any]]]:
    cfg = DOC_CONFIGS[doc.upper()]
    wb = load_workbook(gold_path, data_only=True)
    ws = wb.worksheets[0]
    linked: list[dict[str, Any]] = []
    contract_only: list[dict[str, Any]] = []

    for row in ws.iter_rows(min_row=cfg["start"], values_only=True):
        cid = norm_id(row[cfg["cid"]] if len(row) > cfg["cid"] else "")
        if not cid or cid.upper().startswith("РАЗДЕЛ"):
            continue
        description = str(row[cfg["desc"]] if len(row) > cfg["desc"] and row[cfg["desc"]] else "")
        matrix_ids = parse_ids(row[cfg["analog"]] if len(row) > cfg["analog"] else "")
        status = normalize_status(row[cfg["status"]] if len(row) > cfg["status"] else "", bool(matrix_ids))
        if matrix_ids and status in {"aligned", "deviation"}:
            linked.append(
                {
                    "contract_id": cid,
                    "matrix_ids": sorted(matrix_ids, key=lambda item: numeric_key(item) or (9999,)),
                    "status": status,
                    "description": description,
                }
            )
        elif status == "extra_in_contract" and looks_legal_description(description):
            contract_only.append({"contract_id": cid, "status": status, "description": description})

    matrix_only: list[dict[str, Any]] = []
    if len(wb.worksheets) > 1:
        ws2 = wb.worksheets[1]
        for row in ws2.iter_rows(values_only=True):
            joined = " ".join(str(cell or "") for cell in row[:2])
            for mid in parse_ids(joined):
                description = " ".join(str(cell or "") for cell in row[1:4])
                if not looks_non_legal_description(description):
                    matrix_only.append({"matrix_id": mid, "status": "missing_in_contract"})

    return clean_gold({"linked": linked, "matrix_only": matrix_only, "contract_only": contract_only})


def load_matrix_source_ids(matrix_path: Path | None) -> set[str] | None:
    if not matrix_path or not matrix_path.exists():
        return None
    data = json.loads(matrix_path.read_text(encoding="utf-8-sig"))
    rows: Any
    if isinstance(data, list):
        rows = data
    elif isinstance(data, dict):
        rows = data.get("items") or data.get("matrix") or data.get("data") or data.values()
    else:
        rows = []
    ids = set()
    for row in rows:
        if not isinstance(row, dict):
            continue
        matrix_id = norm_id(row.get("number") or row.get("matrix_id") or row.get("id"))
        if matrix_id:
            ids.add(matrix_id)
    return ids


def clean_gold(gold: dict[str, list[dict[str, Any]]]) -> dict[str, list[dict[str, Any]]]:
    matrix_ids = [row["matrix_id"] for row in gold["matrix_only"]]
    contract_ids = [row["contract_id"] for row in gold["contract_only"]]
    linked_matrix_ids = {
        matrix_id
        for row in gold["linked"]
        for matrix_id in row.get("matrix_ids", [])
    }

    excluded_matrix = {
        mid
        for mid in matrix_ids
        if len(numeric_key(mid) or ()) == 1 or any(other != mid and is_prefix(mid, other) for other in matrix_ids)
    }
    excluded_contract = {
        cid for cid in contract_ids if any(other != cid and is_prefix(cid, other) for other in contract_ids)
    }

    cleaned_matrix_only = [
        row
        for row in gold["matrix_only"]
        if row["matrix_id"] not in excluded_matrix and row["matrix_id"] not in linked_matrix_ids
    ]
    deduped_matrix_only = list({row["matrix_id"]: row for row in cleaned_matrix_only}.values())

    cleaned_contract_only = [
        row for row in gold["contract_only"] if row["contract_id"] not in excluded_contract
    ]
    deduped_contract_only = list({row["contract_id"]: row for row in cleaned_contract_only}.values())

    return {
        "linked": [
            row
            for row in gold["linked"]
            if not looks_non_legal_description(row.get("description", ""))
        ],
        "matrix_only": deduped_matrix_only,
        "contract_only": deduped_contract_only,
    }


def filter_gold_to_matrix_source(
    gold: dict[str, list[dict[str, Any]]],
    source_matrix_ids: set[str] | None,
) -> tuple[dict[str, list[dict[str, Any]]], dict[str, Any]]:
    if not source_matrix_ids:
        return gold, {
            "source_matrix_id_count": 0,
            "source_removed_matrix_only_count": 0,
            "source_affected_linked_rows": 0,
            "source_dropped_linked_rows": 0,
            "source_removed_matrix_ids": [],
        }

    linked: list[dict[str, Any]] = []
    affected = 0
    dropped = 0
    removed_ids: set[str] = set()
    for row in gold["linked"]:
        kept_ids = [matrix_id for matrix_id in row["matrix_ids"] if matrix_id in source_matrix_ids]
        removed = set(row["matrix_ids"]) - set(kept_ids)
        removed_ids.update(removed)
        if removed:
            affected += 1
        if kept_ids:
            new_row = dict(row)
            new_row["matrix_ids"] = kept_ids
            linked.append(new_row)
        else:
            dropped += 1

    matrix_only = []
    removed_matrix_only = 0
    for row in gold["matrix_only"]:
        if row["matrix_id"] in source_matrix_ids:
            matrix_only.append(row)
        else:
            removed_matrix_only += 1
            removed_ids.add(row["matrix_id"])

    return {
        "linked": linked,
        "matrix_only": matrix_only,
        "contract_only": list(gold["contract_only"]),
    }, {
        "source_matrix_id_count": len(source_matrix_ids),
        "source_removed_matrix_only_count": removed_matrix_only,
        "source_affected_linked_rows": affected,
        "source_dropped_linked_rows": dropped,
        "source_removed_matrix_ids": sorted(removed_ids, key=lambda item: numeric_key(item) or (9999,)),
    }


def normalize_agent_status(value: Any) -> str:
    text = str(value or "").strip()
    if text == "full_match":
        return "aligned"
    if text == "partial_match":
        return "deviation"
    if text == "missing":
        return "missing_in_contract"
    return text


def load_agent(artifact_path: Path) -> dict[str, Any]:
    data = json.loads(artifact_path.read_text(encoding="utf-8-sig"))
    groups = []
    for idx, link in enumerate(data.get("links", []) or []):
        groups.append(
            {
                "index": idx,
                "contract_ids": {norm_id(item) for item in link.get("contract_ids", []) or []},
                "matrix_ids": {norm_id(item) for item in link.get("matrix_ids", []) or []},
                "status": normalize_agent_status(link.get("relationship")),
            }
        )
    unmatched_matrix = {
        norm_id(item.get("matrix_id"))
        for item in data.get("unmatched_matrix", []) or []
        if isinstance(item, dict)
    }
    unmatched_contract = {
        norm_id(item.get("contract_id"))
        for item in data.get("unmatched_contract", []) or []
        if isinstance(item, dict) and item.get("status") == "extra_in_contract"
    }
    scope_excluded_matrix = set()
    coverage = data.get("coverage_ledger") or {}
    if isinstance(coverage, dict):
        for item in coverage.get("matrix", []) or []:
            if not isinstance(item, dict):
                continue
            closure = str(item.get("closure") or "").strip()
            reason = str(item.get("reason") or "").lower()
            if closure in {"out_of_scope", "not_applicable"} or (
                closure == "not_evaluable"
                and any(token in reason for token in ("scope", "applic", "filter", "не примен"))
            ):
                scope_excluded_matrix.add(norm_id(item.get("matrix_id")))
    return {
        "groups": groups,
        "unmatched_matrix": unmatched_matrix,
        "unmatched_contract": unmatched_contract,
        "scope_excluded_matrix": scope_excluded_matrix,
    }


def exclude_out_of_scope_gold(
    gold: dict[str, list[dict[str, Any]]],
    scope_excluded: set[str],
) -> tuple[dict[str, list[dict[str, Any]]], dict[str, Any]]:
    if not scope_excluded:
        return gold, {
            "scope_excluded_count": 0,
            "scope_excluded_matrix_ids": [],
            "scope_removed_matrix_only_count": 0,
            "scope_affected_linked_rows": 0,
            "scope_dropped_linked_rows": 0,
        }

    linked: list[dict[str, Any]] = []
    affected = 0
    dropped = 0
    for row in gold["linked"]:
        kept_ids = [matrix_id for matrix_id in row["matrix_ids"] if matrix_id not in scope_excluded]
        if len(kept_ids) != len(row["matrix_ids"]):
            affected += 1
        if kept_ids:
            new_row = dict(row)
            new_row["matrix_ids"] = kept_ids
            linked.append(new_row)
        else:
            dropped += 1

    matrix_only = [row for row in gold["matrix_only"] if row["matrix_id"] not in scope_excluded]
    removed_matrix_only = len(gold["matrix_only"]) - len(matrix_only)
    return {
        "linked": linked,
        "matrix_only": matrix_only,
        "contract_only": list(gold["contract_only"]),
    }, {
        "scope_excluded_count": len(scope_excluded),
        "scope_excluded_matrix_ids": sorted(scope_excluded, key=lambda item: numeric_key(item) or (9999,)),
        "scope_removed_matrix_only_count": removed_matrix_only,
        "scope_affected_linked_rows": affected,
        "scope_dropped_linked_rows": dropped,
    }


def serializable_group(group: dict[str, Any]) -> dict[str, Any]:
    return {
        "index": group.get("index"),
        "contract_ids": sorted(group.get("contract_ids", [])),
        "matrix_ids": sorted(group.get("matrix_ids", [])),
        "status": group.get("status"),
    }


def evaluate(gold: dict[str, list[dict[str, Any]]], agent: dict[str, Any]) -> dict[str, Any]:
    linked_total = len(gold["linked"])
    exact = 0
    soft = 0
    status_hit = 0
    by_status = defaultdict(lambda: {"total": 0, "soft": 0, "status": 0, "exact": 0})
    errors = []

    for row in gold["linked"]:
        gold_contract = row["contract_id"]
        gold_matrix = set(row["matrix_ids"])
        candidates = [
            group
            for group in agent["groups"]
            if gold_contract in group["contract_ids"] and group["matrix_ids"] & gold_matrix
        ]
        exact_candidates = [
            group
            for group in candidates
            if gold_contract in group["contract_ids"] and group["matrix_ids"] == gold_matrix
        ]
        has_soft = bool(candidates)
        has_exact = bool(exact_candidates)
        has_status = any(group["status"] == row["status"] for group in candidates)
        exact += int(has_exact)
        soft += int(has_soft)
        status_hit += int(has_status)
        bucket = by_status[row["status"]]
        bucket["total"] += 1
        bucket["soft"] += int(has_soft)
        bucket["exact"] += int(has_exact)
        bucket["status"] += int(has_status)
        if not has_status:
            errors.append(
                {
                    "type": "status_error" if has_soft else "mapping_miss",
                    "contract_id": gold_contract,
                    "matrix_ids": sorted(gold_matrix),
                    "gold_status": row["status"],
                    "agent_groups": [serializable_group(group) for group in candidates[:3]],
                }
            )

    matrix_total = len(gold["matrix_only"])
    matrix_hit = sum(1 for row in gold["matrix_only"] if row["matrix_id"] in agent["unmatched_matrix"])
    contract_total = len(gold["contract_only"])
    contract_hit = sum(1 for row in gold["contract_only"] if row["contract_id"] in agent["unmatched_contract"])

    total = linked_total + matrix_total + contract_total
    status_total = status_hit + matrix_hit + contract_hit
    soft_total = soft + matrix_hit + contract_hit

    def pct(ok: int, denom: int) -> float:
        return round(ok / denom * 100, 2) if denom else 0.0

    return {
        "gold_counts": {
            "linked": linked_total,
            "matrix_only": matrix_total,
            "contract_only": contract_total,
            "total": total,
        },
        "metrics": {
            "status_accuracy": {"ok": status_total, "total": total, "pct": pct(status_total, total)},
            "linked_mapping_exact_set": {"ok": exact, "total": linked_total, "pct": pct(exact, linked_total)},
            "linked_mapping_soft_1plus": {"ok": soft, "total": linked_total, "pct": pct(soft, linked_total)},
            "linked_group_status": {"ok": status_hit, "total": linked_total, "pct": pct(status_hit, linked_total)},
            "missing_in_contract": {"ok": matrix_hit, "total": matrix_total, "pct": pct(matrix_hit, matrix_total)},
            "extra_in_contract": {"ok": contract_hit, "total": contract_total, "pct": pct(contract_hit, contract_total)},
            "by_status": {
                status: {
                    **bucket,
                    "soft_pct": pct(bucket["soft"], bucket["total"]),
                    "exact_pct": pct(bucket["exact"], bucket["total"]),
                    "status_pct": pct(bucket["status"], bucket["total"]),
                }
                for status, bucket in by_status.items()
            },
        },
        "errors": errors,
    }


def evaluate_with_scope_exclusion(
    gold: dict[str, list[dict[str, Any]]],
    agent: dict[str, Any],
    include_out_of_scope: bool = False,
    source_matrix_ids: set[str] | None = None,
) -> dict[str, Any]:
    gold, source_meta = filter_gold_to_matrix_source(gold, source_matrix_ids)
    if include_out_of_scope:
        result = evaluate(gold, agent)
        result["scope_exclusions"] = {
            "scope_excluded_count": 0,
            "scope_excluded_matrix_ids": [],
            "scope_removed_matrix_only_count": 0,
            "scope_affected_linked_rows": 0,
            "scope_dropped_linked_rows": 0,
        }
        result["source_universe_exclusions"] = source_meta
        return result
    filtered_gold, scope_meta = exclude_out_of_scope_gold(
        gold, agent.get("scope_excluded_matrix", set())
    )
    result = evaluate(filtered_gold, agent)
    result["scope_exclusions"] = scope_meta
    result["source_universe_exclusions"] = source_meta
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--doc", required=True, choices=sorted(DOC_CONFIGS))
    parser.add_argument("--gold", type=Path, required=True)
    parser.add_argument("--artifact", type=Path, required=True)
    parser.add_argument("--matrix-source", type=Path, default=Path("inputs/matrix.json"))
    parser.add_argument("--json-out", type=Path)
    parser.add_argument(
        "--include-out-of-scope",
        action="store_true",
        help="Keep artifact-declared out-of-scope matrix ids in the gold denominator.",
    )
    args = parser.parse_args()

    result = evaluate_with_scope_exclusion(
        load_gold(args.doc, args.gold),
        load_agent(args.artifact),
        include_out_of_scope=args.include_out_of_scope,
        source_matrix_ids=load_matrix_source_ids(args.matrix_source),
    )
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(payload, encoding="utf-8")
    print(payload)


if __name__ == "__main__":
    main()
