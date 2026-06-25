from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import time
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RUNS = ROOT / "runs"
INPUTS = ROOT / "inputs"
ALL_DATA = ROOT / "ALL_DATA"
PYTHON = ROOT / ".venv" / "Scripts" / "python.exe"


@dataclass(frozen=True)
class DocConfig:
    name: str
    txt: Path
    docx: Path
    gold: Path


DOCS = [
    DocConfig(
        "KAVKAZ",
        ALL_DATA / "0.Kavkaz_contract.txt",
        ALL_DATA / "0.Kavkaz.docx",
        ALL_DATA / "KAVKAZ.before_agent_better_pairs_20260618_172816.xlsx",
    ),
    DocConfig("KUZBAS", ALL_DATA / "1.Kyzbas_contract.txt", ALL_DATA / "1.Kyzbas.docx", ALL_DATA / "KUZBAS.xlsx"),
    DocConfig("KALUGA", ALL_DATA / "2.Kaluga_contract.txt", ALL_DATA / "2.Kaluga.docx", ALL_DATA / "KALUGA.xlsx"),
    DocConfig("IRKUTSK", ALL_DATA / "3.Irkutsk_contract.txt", ALL_DATA / "3.Irkutsk.docx", ALL_DATA / "IRKUTSK.xlsx"),
    DocConfig("ALTAI", ALL_DATA / "altai_contract.txt", ALL_DATA / "4.Altai.docx", ALL_DATA / "ALTAI.xlsx"),
]


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def norm_id(value: Any) -> str:
    text = str(value or "").strip()
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"^(п\.?|пункт)\s+", "", text, flags=re.I)
    text = text.strip(" ;,")
    while text.endswith("."):
        text = text[:-1].strip()
    return text


def numeric_key(value: str) -> tuple[int, ...] | None:
    value = norm_id(value)
    if not re.fullmatch(r"\d+(?:\.\d+)*", value):
        return None
    return tuple(int(part) for part in value.split("."))


def is_prefix_id(parent: str, child: str) -> bool:
    parent = norm_id(parent)
    child = norm_id(child)
    return bool(parent and child and child.startswith(parent + "."))


def parse_matrix_ids(value: Any) -> set[str]:
    text = str(value or "")
    if not text or text.strip() in {"—", "-", "0"}:
        return set()
    return {norm_id(match.group(0)) for match in re.finditer(r"\d+(?:\.\d+)+\.?", text)}


def parse_contract_locator(value: Any) -> str:
    return norm_id(value)


def empty_analog(value: Any) -> bool:
    text = str(value or "").strip().lower()
    return not text or text in {"—", "-", "0", "нет", "нет аналога"}


def normalize_gold_status(value: Any, has_matrix_analogs: bool) -> str | None:
    text = str(value or "").lower()
    if "расхожд" in text or "partial" in text or "deviation" in text:
        return "deviation"
    if "соответ" in text or "full" in text or "aligned" in text:
        return "aligned"
    if "нет аналога" in text:
        if "матриц" in text or not has_matrix_analogs:
            return "extra_in_contract"
        return "missing_in_contract"
    if has_matrix_analogs:
        return "aligned"
    return None


def normalize_agent_status(value: Any) -> str:
    text = str(value or "").strip()
    if text == "full_match":
        return "aligned"
    if text == "partial_match":
        return "deviation"
    if text == "missing":
        return "missing_in_contract"
    return text


def find_header(ws, required: list[str]) -> tuple[int, dict[str, int]]:
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        lowered = [str(cell or "").lower() for cell in row]
        if all(any(req in cell for cell in lowered) for req in required):
            mapping = {str(cell or "").lower(): idx for idx, cell in enumerate(row)}
            return row_idx, mapping
    raise RuntimeError(f"Header not found in {ws.title!r}: {required}")


def detect_first_sheet_columns(ws) -> tuple[int, int, int, int]:
    header_row = None
    header = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        lowered = [str(cell or "").lower() for cell in row]
        if any("пункт" in c and ("договор" in c or "контракт" in c or "п/п" in c) for c in lowered) and any(
            "аналог" in c and "матриц" in c for c in lowered
        ) and any("статус" in c for c in lowered):
            header_row = row_idx
            header = lowered
            break
    if header_row is None:
        raise RuntimeError(f"Could not detect first-sheet header in {ws.title}")

    contract_col = next(
        i
        for i, c in enumerate(header)
        if "пункт" in c and ("договор" in c or "контракт" in c or "п/п" in c)
    )
    matrix_col = next(i for i, c in enumerate(header) if "аналог" in c and "матриц" in c)
    status_col = next(i for i, c in enumerate(header) if "статус" in c)
    return header_row, contract_col, matrix_col, status_col


def detect_matrix_only_columns(ws) -> tuple[int, int]:
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        lowered = [str(cell or "").lower() for cell in row]
        if any("пункт" in c and "матриц" in c for c in lowered):
            matrix_col = next(i for i, c in enumerate(lowered) if "пункт" in c and "матриц" in c)
            return row_idx, matrix_col
    raise RuntimeError(f"Could not detect matrix-only header in {ws.title}")


def load_gold(gold_path: Path) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(gold_path, read_only=True, data_only=True)
    first = wb.worksheets[0]
    matrix_only = wb.worksheets[1] if len(wb.worksheets) > 1 else None

    header_row, contract_col, matrix_col, status_col = detect_first_sheet_columns(first)
    linked: list[dict[str, Any]] = []
    contract_only: list[dict[str, Any]] = []

    for row in first.iter_rows(min_row=header_row + 1, values_only=True):
        contract_id = parse_contract_locator(row[contract_col] if contract_col < len(row) else "")
        if not contract_id or "раздел" in contract_id.lower():
            continue
        matrix_ids = parse_matrix_ids(row[matrix_col] if matrix_col < len(row) else "")
        status = normalize_gold_status(row[status_col] if status_col < len(row) else "", bool(matrix_ids))
        if matrix_ids:
            if status in {"aligned", "deviation"}:
                linked.append({"contract_id": contract_id, "matrix_ids": matrix_ids, "status": status})
        elif status == "extra_in_contract" or empty_analog(row[matrix_col] if matrix_col < len(row) else ""):
            contract_only.append({"contract_id": contract_id, "status": "extra_in_contract"})

    matrix_missing: list[dict[str, Any]] = []
    if matrix_only is not None:
        header_row2, matrix_col2 = detect_matrix_only_columns(matrix_only)
        for row in matrix_only.iter_rows(min_row=header_row2 + 1, values_only=True):
            cell = row[matrix_col2] if matrix_col2 < len(row) else ""
            ids = parse_matrix_ids(cell)
            if not ids:
                direct = norm_id(cell)
                if re.fullmatch(r"\d+(?:\.\d+)*", direct):
                    ids = {direct}
            for mid in ids:
                matrix_missing.append({"matrix_id": mid, "status": "missing_in_contract"})

    return {"linked": linked, "matrix_only": matrix_missing, "contract_only": contract_only}


def clean_gold(gold: dict[str, list[dict[str, Any]]]) -> tuple[dict[str, list[dict[str, Any]]], dict[str, list[str]]]:
    matrix_ids = [row["matrix_id"] for row in gold["matrix_only"]]
    contract_ids = [row["contract_id"] for row in gold["contract_only"]]

    excluded_matrix = sorted(
        {
            mid
            for mid in matrix_ids
            if len(numeric_key(mid) or ()) == 1 or any(is_prefix_id(mid, other) for other in matrix_ids)
        }
    )
    excluded_contract = sorted({cid for cid in contract_ids if any(is_prefix_id(cid, other) for other in contract_ids)})

    clean = {
        "linked": list(gold["linked"]),
        "matrix_only": [row for row in gold["matrix_only"] if row["matrix_id"] not in excluded_matrix],
        "contract_only": [row for row in gold["contract_only"] if row["contract_id"] not in excluded_contract],
    }
    return clean, {"matrix_only": excluded_matrix, "contract_only": excluded_contract}


def id_matches(gold_locator: str, agent_locator: str) -> bool:
    gold_locator = norm_id(gold_locator)
    agent_locator = norm_id(agent_locator)
    if gold_locator == agent_locator:
        return True
    if "-" in gold_locator:
        left, right = [norm_id(part) for part in gold_locator.split("-", 1)]
        lk = numeric_key(left)
        rk = numeric_key(right)
        ak = numeric_key(agent_locator)
        if lk and rk and ak and len(lk) == len(rk) == len(ak):
            return lk <= ak <= rk
    return False


def load_agent(artifact_path: Path) -> dict[str, Any]:
    artifact = json.loads(artifact_path.read_text(encoding="utf-8-sig"))
    atoms = []
    for atom in artifact.get("atomic_links", []) or []:
        if not isinstance(atom, dict):
            continue
        atoms.append(
            {
                "matrix_id": norm_id(atom.get("matrix_id")),
                "contract_id": norm_id(atom.get("contract_id")),
                "status": normalize_agent_status(atom.get("relationship")),
            }
        )
    for link in artifact.get("links", []) or []:
        status = normalize_agent_status(link.get("relationship"))
        for mid in link.get("matrix_ids", []) or []:
            for cid in link.get("contract_ids", []) or []:
                pair = {"matrix_id": norm_id(mid), "contract_id": norm_id(cid), "status": status}
                if pair not in atoms:
                    atoms.append(pair)

    unmatched_matrix = {norm_id(row.get("matrix_id")) for row in artifact.get("unmatched_matrix", []) or [] if isinstance(row, dict)}
    unmatched_contract = [
        {"contract_id": norm_id(row.get("contract_id")), "status": normalize_agent_status(row.get("status"))}
        for row in artifact.get("unmatched_contract", []) or []
        if isinstance(row, dict)
    ]
    return {
        "atoms": atoms,
        "unmatched_matrix": unmatched_matrix,
        "unmatched_contract": unmatched_contract,
        "summary": artifact.get("summary", {}),
        "counts": {
            "links": len(artifact.get("links", []) or []),
            "atomic_links": len(artifact.get("atomic_links", []) or []),
            "unmatched_matrix": len(artifact.get("unmatched_matrix", []) or []),
            "unmatched_contract": len(artifact.get("unmatched_contract", []) or []),
        },
    }


def evaluate_gold(gold: dict[str, list[dict[str, Any]]], agent: dict[str, Any]) -> dict[str, Any]:
    totals = {
        "linked_total": len(gold["linked"]),
        "linked_mapping_hit": 0,
        "linked_status_hit": 0,
        "matrix_only_total": len(gold["matrix_only"]),
        "matrix_only_hit": 0,
        "contract_only_total": len(gold["contract_only"]),
        "contract_only_hit": 0,
    }
    by_status: dict[str, dict[str, int]] = {}

    def bump(status: str, mapping: bool, status_ok: bool) -> None:
        bucket = by_status.setdefault(status, {"total": 0, "mapping_hit": 0, "status_hit": 0})
        bucket["total"] += 1
        bucket["mapping_hit"] += int(mapping)
        bucket["status_hit"] += int(status_ok)

    atoms = agent["atoms"]
    for row in gold["linked"]:
        matching = [
            atom
            for atom in atoms
            if id_matches(row["contract_id"], atom["contract_id"]) and atom["matrix_id"] in row["matrix_ids"]
        ]
        mapping = bool(matching)
        if row["status"] == "deviation":
            status_ok = any(atom["status"] == "deviation" for atom in matching)
        else:
            status_ok = any(atom["status"] == "aligned" for atom in matching)
        totals["linked_mapping_hit"] += int(mapping)
        totals["linked_status_hit"] += int(status_ok)
        bump(row["status"], mapping, status_ok)

    for row in gold["matrix_only"]:
        linked_to_matrix = any(atom["matrix_id"] == row["matrix_id"] for atom in atoms)
        mapping = row["matrix_id"] in agent["unmatched_matrix"] or not linked_to_matrix
        status_ok = row["matrix_id"] in agent["unmatched_matrix"]
        totals["matrix_only_hit"] += int(status_ok)
        bump("missing_in_contract", mapping, status_ok)

    for row in gold["contract_only"]:
        hits = [item for item in agent["unmatched_contract"] if id_matches(row["contract_id"], item["contract_id"])]
        mapping = bool(hits)
        status_ok = any(item["status"] == "extra_in_contract" for item in hits)
        totals["contract_only_hit"] += int(status_ok)
        bump("extra_in_contract", mapping, status_ok)

    total = totals["linked_total"] + totals["matrix_only_total"] + totals["contract_only_total"]
    mapping_hit = totals["linked_mapping_hit"] + totals["matrix_only_hit"] + totals["contract_only_hit"]
    status_hit = totals["linked_status_hit"] + totals["matrix_only_hit"] + totals["contract_only_hit"]

    result = {
        **totals,
        "total": total,
        "mapping_hit": mapping_hit,
        "status_hit": status_hit,
        "mapping_accuracy": pct(mapping_hit, total),
        "status_accuracy": pct(status_hit, total),
        "linked_mapping_accuracy": pct(totals["linked_mapping_hit"], totals["linked_total"]),
        "linked_status_accuracy": pct(totals["linked_status_hit"], totals["linked_total"]),
        "status_on_mapped_linked_accuracy": pct(totals["linked_status_hit"], totals["linked_mapping_hit"]),
        "by_status": {},
    }
    for status, bucket in by_status.items():
        result["by_status"][status] = {
            **bucket,
            "mapping_accuracy": pct(bucket["mapping_hit"], bucket["total"]),
            "status_accuracy": pct(bucket["status_hit"], bucket["total"]),
        }
    return result


def pct(num: int, den: int) -> float:
    return round(num * 100 / den, 2) if den else 0.0


def evaluate_artifact(artifact_path: Path, gold_path: Path) -> dict[str, Any]:
    gold = load_gold(gold_path)
    clean, exclusions = clean_gold(gold)
    agent = load_agent(artifact_path)
    return {
        "raw": evaluate_gold(gold, agent),
        "clean": evaluate_gold(clean, agent),
        "agent_counts": agent["counts"],
        "artifact_summary": agent["summary"],
        "clean_exclusions": exclusions,
    }


def existing_numbered_score(text: str) -> tuple[int, int]:
    lines = text.splitlines()
    numbered = [line for line in lines if re.search(r"^\s*(?:\d+(?:\.\d+)+|\d+)\.?\s+", line)]
    deep = [line for line in lines if re.search(r"^\s*\d+\.\d+(?:\.\d+)*\.?\s+", line)]
    return len(numbered), len(deep)


def paragraph_text(p: ET.Element) -> str:
    parts: list[str] = []
    for node in p.iter():
        tag = node.tag
        if tag == f"{{{NS['w']}}}t" and node.text:
            parts.append(node.text)
        elif tag == f"{{{NS['w']}}}tab":
            parts.append("\t")
        elif tag == f"{{{NS['w']}}}br":
            parts.append("\n")
    return "".join(parts).strip()


def read_numbering(zf: zipfile.ZipFile) -> tuple[dict[str, str], dict[str, dict[int, dict[str, str]]]]:
    try:
        root = ET.fromstring(zf.read("word/numbering.xml"))
    except KeyError:
        return {}, {}
    abstract: dict[str, dict[int, dict[str, str]]] = {}
    for abs_num in root.findall("w:abstractNum", NS):
        abs_id = abs_num.attrib.get(f"{{{NS['w']}}}abstractNumId", "")
        levels: dict[int, dict[str, str]] = {}
        for lvl in abs_num.findall("w:lvl", NS):
            ilvl = int(lvl.attrib.get(f"{{{NS['w']}}}ilvl", "0"))
            fmt = lvl.find("w:numFmt", NS)
            text = lvl.find("w:lvlText", NS)
            start = lvl.find("w:start", NS)
            levels[ilvl] = {
                "fmt": fmt.attrib.get(f"{{{NS['w']}}}val", "decimal") if fmt is not None else "decimal",
                "text": text.attrib.get(f"{{{NS['w']}}}val", f"%{ilvl + 1}.") if text is not None else f"%{ilvl + 1}.",
                "start": start.attrib.get(f"{{{NS['w']}}}val", "1") if start is not None else "1",
            }
        abstract[abs_id] = levels

    num_to_abs: dict[str, str] = {}
    for num in root.findall("w:num", NS):
        num_id = num.attrib.get(f"{{{NS['w']}}}numId", "")
        abs_el = num.find("w:abstractNumId", NS)
        if abs_el is not None:
            num_to_abs[num_id] = abs_el.attrib.get(f"{{{NS['w']}}}val", "")
    return num_to_abs, abstract


def num_pr(p: ET.Element) -> tuple[str, int] | None:
    ppr = p.find("w:pPr", NS)
    if ppr is None:
        return None
    npr = ppr.find("w:numPr", NS)
    if npr is None:
        return None
    num_id_el = npr.find("w:numId", NS)
    ilvl_el = npr.find("w:ilvl", NS)
    if num_id_el is None:
        return None
    num_id = num_id_el.attrib.get(f"{{{NS['w']}}}val", "")
    ilvl = int(ilvl_el.attrib.get(f"{{{NS['w']}}}val", "0")) if ilvl_el is not None else 0
    return num_id, ilvl


def format_label(lvl_text: str, counters: dict[int, int]) -> str:
    label = lvl_text
    for idx in range(1, 10):
        label = label.replace(f"%{idx}", str(counters.get(idx - 1, 0)))
    return label


def extract_docx_numbered(docx_path: Path) -> str:
    with zipfile.ZipFile(docx_path) as zf:
        document = ET.fromstring(zf.read("word/document.xml"))
        num_to_abs, abstract = read_numbering(zf)

    counters_by_num: dict[str, dict[int, int]] = {}
    lines: list[str] = []
    for p in document.iter(f"{{{NS['w']}}}p"):
        text = paragraph_text(p)
        if not text:
            continue
        label = ""
        np = num_pr(p)
        if np:
            num_id, ilvl = np
            abs_id = num_to_abs.get(num_id, "")
            level = abstract.get(abs_id, {}).get(ilvl)
            if level and level.get("fmt") == "decimal":
                counters = counters_by_num.setdefault(num_id, {})
                levels = abstract.get(abs_id, {})
                for parent in range(ilvl):
                    if parent not in counters:
                        parent_start = int(levels.get(parent, {}).get("start") or "1")
                        counters[parent] = parent_start
                start = int(level.get("start") or "1")
                counters[ilvl] = counters.get(ilvl, start - 1) + 1
                for deeper in list(counters):
                    if deeper > ilvl:
                        del counters[deeper]
                label = format_label(level.get("text") or f"%{ilvl + 1}.", counters)
        if label and not re.match(r"^\s*\d+(?:\.\d+)*\.?\s+", text):
            line = f"{label} {text}"
        else:
            line = text
        lines.append(re.sub(r"[ \t]+", " ", line).strip())
    return "\n".join(lines) + "\n"


def prepare_contract_text(doc: DocConfig, run_dir: Path) -> dict[str, Any]:
    source = doc.txt
    text = source.read_text(encoding="utf-8-sig", errors="replace") if source.exists() else ""
    numbered, deep = existing_numbered_score(text)
    extracted = False
    if numbered < 50 or deep < 30:
        text = extract_docx_numbered(doc.docx)
        extracted = True
        extracted_path = run_dir / f"{doc.name.lower()}_contract_numbered.txt"
        extracted_path.write_text(text, encoding="utf-8")
        numbered, deep = existing_numbered_score(text)
        source = extracted_path
    shutil.copyfile(source, INPUTS / "contract.txt")
    return {
        "source": str(source.relative_to(ROOT) if source.is_relative_to(ROOT) else source),
        "extracted_from_docx": extracted,
        "numbered_lines": numbered,
        "deep_numbered_lines": deep,
        "bytes": len(text.encode("utf-8")),
    }


def run_one(doc: DocConfig, stamp: str) -> dict[str, Any]:
    run_dir = RUNS / f"all_docs_{stamp}" / doc.name.lower()
    run_dir.mkdir(parents=True, exist_ok=True)
    prep = prepare_contract_text(doc, run_dir)
    stdout = run_dir / "stdout.log"
    stderr = run_dir / "stderr.log"
    started = time.time()
    with stdout.open("w", encoding="utf-8", errors="replace") as out, stderr.open("w", encoding="utf-8", errors="replace") as err:
        proc = subprocess.run(
            [str(PYTHON), "-u", "main.py"],
            cwd=ROOT,
            stdout=out,
            stderr=err,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
    elapsed = round(time.time() - started, 2)
    artifact_src = ROOT / "outputs" / "discrepancy_analysis.json"
    artifact_dst = run_dir / "discrepancy_analysis.json"
    metrics = None
    if artifact_src.exists():
        shutil.copyfile(artifact_src, artifact_dst)
        metrics = evaluate_artifact(artifact_dst, doc.gold)
        (run_dir / "metrics.json").write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    result = {
        "doc": doc.name,
        "returncode": proc.returncode,
        "elapsed_sec": elapsed,
        "contract_text": prep,
        "artifact": str(artifact_dst.relative_to(ROOT)) if artifact_dst.exists() else None,
        "metrics": metrics,
    }
    (run_dir / "run_result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False)[:4000], flush=True)
    return result


def main() -> None:
    RUNS.mkdir(exist_ok=True)
    INPUTS.mkdir(exist_ok=True)
    if not PYTHON.exists():
        raise SystemExit(f"Missing venv python: {PYTHON}")
    selected = [arg.upper() for arg in sys.argv[1:]]
    docs = DOCS
    if selected:
        docs = [doc for doc in DOCS if doc.name.upper() in selected]
        known = {doc.name.upper() for doc in docs}
        unknown = [name for name in selected if name not in known]
        if unknown:
            raise SystemExit(f"Unknown doc name(s): {', '.join(unknown)}")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    summary = []
    for doc in docs:
        print(f"=== RUN {doc.name} ===", flush=True)
        summary.append(run_one(doc, stamp))
    out_dir = RUNS / f"all_docs_{stamp}"
    (out_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"SUMMARY_PATH={out_dir / 'summary.json'}", flush=True)


if __name__ == "__main__":
    main()
