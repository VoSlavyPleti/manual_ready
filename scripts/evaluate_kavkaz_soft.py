from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_GOLD = ROOT / "ALL_DATA" / "KAVKAZ.before_agent_better_pairs_20260618_172816.xlsx"
DEFAULT_AGENT = ROOT / "outputs" / "matrix_contract_mapping.json"
DEFAULT_MATRIX = ROOT / "inputs" / "matrix.json"
DEFAULT_OUT = ROOT / "outputs" / "cold_start_kavkaz_soft_metrics.json"


STATUS_MAP = {
    "✅ Соответствует": "full_match",
    "⚠️ Расхождение": "partial_match",
    "❌ Нет аналога": "missing",
    "❌ Нет аналога в контракте": "missing",
}


def maybe_fix_mojibake(value: str) -> str:
    try:
        fixed = value.encode("cp1251").decode("utf-8")
    except UnicodeError:
        return value
    return fixed if fixed.count("�") <= value.count("�") else value


def normalize_id(value: Any) -> str:
    if value is None:
        return ""
    text = maybe_fix_mojibake(str(value))
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    text = text.strip(" ,;")

    leading_clause = re.match(
        r"^п\.?\s*(\d+(?:\.\d+)*\.?)(?:\s+(?:абз\.?|абзац)\s*\d+)?$",
        text,
        re.I,
    )
    if leading_clause:
        text = leading_clause.group(1)

    numeric = re.match(r"^(\d+(?:\.\d+)*\.?)\s*\(", text)
    if numeric:
        text = numeric.group(1)

    appendix = re.search(r"приложени[ея]\s*№?\s*([0-9]+(?:\.[0-9]+)*)", text, re.I)
    appendix_point = re.search(r"\bп\.?\s*([0-9]+(?:\.[0-9]+)*)", text, re.I)
    if appendix:
        text = f"Приложение №{appendix.group(1)}"
        if appendix_point:
            text += f" п.{appendix_point.group(1)}"

    text = re.sub(r"(?<=\d)\.$", "", text)
    text = re.sub(r"\s*№\s*", " №", text)
    text = text.replace("Приложение N", "Приложение №")
    return text.lower()


def display_id(value: Any) -> str:
    if value is None:
        return ""
    text = maybe_fix_mojibake(str(value)).replace("\u00a0", " ")
    return re.sub(r"\s+", " ", text).strip().strip(" ,;")


def split_ids(value: Any) -> list[str]:
    if value is None:
        return []
    text = display_id(value)
    if not text:
        return []
    parts = re.split(r"[,;\n]+", text)
    return [display_id(part) for part in parts if display_id(part)]


def normalize_status(value: Any) -> str:
    if value is None:
        return ""
    text = display_id(value)
    if text in STATUS_MAP:
        return STATUS_MAP[text]
    lower = text.lower()
    if "full" in lower or "соответ" in lower:
        return "full_match"
    if "partial" in lower or "расхожд" in lower:
        return "partial_match"
    if "missing" in lower or "нет аналога" in lower:
        return "missing"
    return lower


def load_agent(path: Path) -> dict[str, dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    rows: dict[str, dict[str, Any]] = {}
    for row in data:
        matrix_id = normalize_id(row.get("matrix_id"))
        analogs = row.get("contract_analog") or []
        if isinstance(analogs, str):
            analogs = split_ids(analogs)
        rows[matrix_id] = {
            "raw": row,
            "status": normalize_status(row.get("overall_status")),
            "contract_ids": {normalize_id(item) for item in analogs if normalize_id(item)},
            "contract_display": [display_id(item) for item in analogs if display_id(item)],
        }
    return rows


def load_matrix_ids(path: Path) -> set[str]:
    matrix = json.loads(path.read_text(encoding="utf-8-sig"))
    return {normalize_id(item.get("number")) for item in matrix if normalize_id(item.get("number"))}


def load_gold(path: Path, matrix_scope: set[str]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    wb = load_workbook(path, data_only=True)
    pair_ws = wb["🔁 Контракт ↔ Матрица"]
    missing_ws = wb["Только в матрице"]

    pair_rows: list[dict[str, Any]] = []
    skipped = Counter()
    for row in pair_ws.iter_rows(min_row=3, values_only=True):
        contract_id, _summary, status, matrix_analogs, comment = row[:5]
        contract_display = display_id(contract_id)
        analogs = split_ids(matrix_analogs)
        scoped_analogs = [mid for mid in analogs if normalize_id(mid) in matrix_scope]
        if not contract_display:
            skipped["empty_contract_id"] += 1
            continue
        if not analogs:
            skipped["contract_only_no_matrix_analog"] += 1
            continue
        if not scoped_analogs:
            skipped["analogs_outside_current_matrix_scope"] += 1
            continue
        pair_rows.append(
            {
                "type": "contract_row",
                "contract_id": contract_display,
                "contract_norm": normalize_id(contract_display),
                "matrix_ids": scoped_analogs,
                "matrix_norms": [normalize_id(mid) for mid in scoped_analogs],
                "gold_status": normalize_status(status),
                "gold_status_raw": display_id(status),
                "comment": display_id(comment),
            }
        )

    missing_rows: list[dict[str, Any]] = []
    for row in missing_ws.iter_rows(min_row=3, values_only=True):
        matrix_id, _summary, status, note, *_rest = row
        matrix_display = display_id(matrix_id)
        matrix_norm = normalize_id(matrix_display)
        if not matrix_norm:
            skipped["matrix_only_empty_id"] += 1
            continue
        if matrix_norm not in matrix_scope:
            skipped["matrix_only_outside_current_scope"] += 1
            continue
        if normalize_status(status) != "missing":
            skipped["matrix_only_non_missing_status"] += 1
            continue
        missing_rows.append(
            {
                "type": "matrix_only",
                "matrix_id": matrix_display,
                "matrix_norm": matrix_norm,
                "gold_status": "missing",
                "note": display_id(note),
            }
        )

    return pair_rows, missing_rows, dict(skipped)


def evaluate(agent: dict[str, dict[str, Any]], pair_rows: list[dict[str, Any]], missing_rows: list[dict[str, Any]]) -> dict[str, Any]:
    evaluated: list[dict[str, Any]] = []
    errors_by_type: dict[str, list[dict[str, Any]]] = defaultdict(list)
    status_counts = Counter()
    hit_counts = Counter()

    for gold in pair_rows:
        matched_matrix_ids = []
        matched_status_ids = []
        agent_candidates = {}
        agent_statuses = {}
        for matrix_norm, matrix_display in zip(gold["matrix_norms"], gold["matrix_ids"]):
            row = agent.get(matrix_norm)
            if not row:
                continue
            agent_candidates[matrix_display] = row["contract_display"]
            agent_statuses[matrix_display] = row["status"]
            if gold["contract_norm"] in row["contract_ids"]:
                matched_matrix_ids.append(matrix_display)
                if row["status"] == gold["gold_status"]:
                    matched_status_ids.append(matrix_display)

        mapping_hit = bool(matched_matrix_ids)
        status_hit = bool(matched_status_ids)
        status_counts[gold["gold_status"]] += 1
        hit_counts["total"] += 1
        hit_counts["mapping_hit"] += int(mapping_hit)
        hit_counts["status_hit"] += int(status_hit)
        hit_counts[f"status_total::{gold['gold_status']}"] += 1
        hit_counts[f"status_hit::{gold['gold_status']}"] += int(status_hit)
        hit_counts[f"mapping_hit::{gold['gold_status']}"] += int(mapping_hit)

        item = {
            **gold,
            "mapping_hit": mapping_hit,
            "status_hit": status_hit,
            "matched_matrix_ids": matched_matrix_ids,
            "matched_status_ids": matched_status_ids,
            "agent_candidates_by_matrix": agent_candidates,
            "agent_statuses_by_matrix": agent_statuses,
        }
        evaluated.append(item)

        if not mapping_hit:
            any_agent_contracts = any(agent.get(mid, {}).get("contract_ids") for mid in gold["matrix_norms"])
            error_type = "wrong_candidate_only" if any_agent_contracts else "false_missing"
            errors_by_type[error_type].append(item)
        elif not status_hit:
            errors_by_type["status_error_after_mapping_hit"].append(item)

    for gold in missing_rows:
        row = agent.get(gold["matrix_norm"])
        contract_ids = row["contract_display"] if row else []
        agent_status = row["status"] if row else ""
        mapping_hit = row is not None and not contract_ids
        status_hit = mapping_hit and agent_status == "missing"
        status_counts["missing"] += 1
        hit_counts["total"] += 1
        hit_counts["mapping_hit"] += int(mapping_hit)
        hit_counts["status_hit"] += int(status_hit)
        hit_counts["status_total::missing"] += 1
        hit_counts["status_hit::missing"] += int(status_hit)
        hit_counts["mapping_hit::missing"] += int(mapping_hit)

        item = {
            **gold,
            "mapping_hit": mapping_hit,
            "status_hit": status_hit,
            "agent_contract_ids": contract_ids,
            "agent_status": agent_status,
        }
        evaluated.append(item)
        if contract_ids:
            errors_by_type["false_non_missing"].append(item)
        elif not status_hit:
            errors_by_type["missing_status_error"].append(item)

    metrics = {
        "total": hit_counts["total"],
        "mapping_hit": hit_counts["mapping_hit"],
        "status_hit": hit_counts["status_hit"],
        "mapping_accuracy": hit_counts["mapping_hit"] / hit_counts["total"] if hit_counts["total"] else 0,
        "soft_status_accuracy": hit_counts["status_hit"] / hit_counts["total"] if hit_counts["total"] else 0,
        "by_gold_status": {},
        "gold_status_distribution": dict(status_counts),
    }
    for status in sorted(status_counts):
        total = hit_counts[f"status_total::{status}"]
        metrics["by_gold_status"][status] = {
            "total": total,
            "mapping_hit": hit_counts[f"mapping_hit::{status}"],
            "status_hit": hit_counts[f"status_hit::{status}"],
            "status_accuracy": hit_counts[f"status_hit::{status}"] / total if total else 0,
        }

    return {
        "metrics": metrics,
        "errors": {key: value for key, value in errors_by_type.items()},
        "error_counts": {key: len(value) for key, value in errors_by_type.items()},
        "evaluated_rows": evaluated,
    }


def main() -> None:
    matrix_scope = load_matrix_ids(DEFAULT_MATRIX)
    agent = load_agent(DEFAULT_AGENT)
    pair_rows, missing_rows, skipped = load_gold(DEFAULT_GOLD, matrix_scope)
    result = evaluate(agent, pair_rows, missing_rows)
    result["gold_rows"] = {
        "contract_rows": len(pair_rows),
        "matrix_only_rows": len(missing_rows),
        "skipped": skipped,
    }
    DEFAULT_OUT.parent.mkdir(parents=True, exist_ok=True)
    DEFAULT_OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    metrics = result["metrics"]
    print(f"total={metrics['total']}")
    print(f"mapping_hit={metrics['mapping_hit']} ({metrics['mapping_accuracy']:.2%})")
    print(f"soft_status_hit={metrics['status_hit']} ({metrics['soft_status_accuracy']:.2%})")
    print("by_gold_status=")
    for status, row in metrics["by_gold_status"].items():
        print(
            f"  {status}: {row['status_hit']}/{row['total']} "
            f"({row['status_accuracy']:.2%}), mapping {row['mapping_hit']}/{row['total']}"
        )
    print("error_counts=", result["error_counts"])
    print("gold_rows=", result["gold_rows"])
    print(f"saved={DEFAULT_OUT}")


if __name__ == "__main__":
    main()
