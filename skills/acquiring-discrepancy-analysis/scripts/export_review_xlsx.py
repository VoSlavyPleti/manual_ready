"""Export discrepancy-analysis JSON to a two-sheet review workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


STATUS_LABELS = {
    "aligned": "aligned",
    "deviation": "deviation",
    "missing_in_contract": "missing_in_contract",
    "extra_in_contract": "extra_in_contract",
}


def join_ids(values: list[str]) -> str:
    return ", ".join(str(item) for item in values if str(item).strip())


def export(artifact_path: Path, xlsx_path: Path) -> None:
    artifact = json.loads(artifact_path.read_text(encoding="utf-8-sig"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Contract-Matrix"
    ws.append(["contract_id", "contract_summary", "matrix_ids", "status", "risk_level", "comment"])

    for link in artifact.get("links", []) or []:
        status = STATUS_LABELS.get(link.get("relationship"), link.get("relationship", ""))
        comment = link.get("status_reason") or ""
        if link.get("discrepancies"):
            gaps = []
            for item in link.get("discrepancies", []) or []:
                if isinstance(item, dict):
                    gaps.append(item.get("description") or item.get("risk") or "")
            comment = "; ".join([comment, *[gap for gap in gaps if gap]]).strip("; ")
        ws.append(
            [
                join_ids(link.get("contract_ids", []) or []),
                link.get("contract_position", ""),
                join_ids(link.get("matrix_ids", []) or []),
                status,
                link.get("risk_level", ""),
                comment,
            ]
        )

    for row in artifact.get("unmatched_contract", []) or []:
        if row.get("status") != "extra_in_contract":
            continue
        ws.append(
            [
                row.get("contract_id", ""),
                row.get("contract_position", ""),
                "",
                "extra_in_contract",
                row.get("risk_level", ""),
                row.get("risk") or row.get("materiality_reason", ""),
            ]
        )

    ws2 = wb.create_sheet("Only in Matrix")
    ws2.append(["matrix_id", "requirement", "required_type", "risk_level", "comment"])
    for row in artifact.get("unmatched_matrix", []) or []:
        ws2.append(
            [
                row.get("matrix_id", ""),
                row.get("requirement", ""),
                row.get("required_type", ""),
                row.get("risk_level", ""),
                row.get("risk", ""),
            ]
        )

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            width = min(70, max(12, max(len(str(cell.value or "")) for cell in sheet[letter]) + 2))
            sheet.column_dimensions[letter].width = width
        sheet.freeze_panes = "A2"

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("artifact", type=Path)
    parser.add_argument("xlsx_out", type=Path)
    args = parser.parse_args()
    export(args.artifact, args.xlsx_out)
    print(f"wrote {args.xlsx_out}")


if __name__ == "__main__":
    main()
